"""PDF (WeasyPrint, embedded Ethiopic font via fontconfig) and Excel (openpyxl) exports.

Depends on exact field/attribute names from models.py — see the docstring on each
function for the precise shape it expects, since these functions get called with
plain dicts/objects assembled by the views, not directly from the ORM.
"""
from io import BytesIO

from django.conf import settings
from django.http import FileResponse
from django.template.loader import render_to_string

from . import ethiopic

# AttendanceEntry.status uses the full TextChoices words, not single-letter codes.
STATUS_SHORT = {"present": "ተ", "absent": "ቀ", "permission": "ፈ", "": ""}


def render_pdf_bytes(template, context):
    from weasyprint import HTML  # imported lazily: needs system libraries (see README)

    html = render_to_string(template, context)
    return HTML(string=html, base_url=str(settings.BASE_DIR)).write_pdf()


def pdf_response(template, context, filename):
    data = render_pdf_bytes(template, context)
    return FileResponse(BytesIO(data), as_attachment=True, filename=filename, content_type="application/pdf")


def xlsx_response(workbook, filename):
    buf = BytesIO()
    workbook.save(buf)
    buf.seek(0)
    return FileResponse(
        buf, as_attachment=True, filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _style_header(ws, row, ncols):
    from openpyxl.styles import Alignment, Font, PatternFill

    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE7F3")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def sheet_title(sheet):
    """AttendanceSheet has no stored title — compute the Amharic month/year label from
    ethiopian_year/ethiopian_month, or mark it as the template. Views and templates
    needing a display label for a sheet should call this rather than inventing their own."""
    if sheet.is_template:
        return "አብነት (Template)"
    return f"{ethiopic.MONTHS[sheet.ethiopian_month - 1]} {sheet.ethiopian_year}"


def attendance_workbook(sheet, sundays, rows):
    """`sundays`: list of Gregorian dates (e.g. from ethiopic.sundays_in_month).
    `rows`: list of dicts — {"n", "student" (has .full_name), "cells": [{"date", "status"}],
    "present", "absent", "permission"} — this is the shape the attendance view must build
    from AttendanceRow/AttendanceEntry before calling this function."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "መገኘት"[:31]
    ws.append([f"{sheet.classroom.name} — {sheet_title(sheet)}"])
    ws.append([])
    columns = ["#", "ስም (Name)"]
    if sheet.is_template:
        columns += [f"ሳምንት {i}" for i in range(1, 6)]
        span = 5
    else:
        columns += [ethiopic.format_short(d) for d in sundays]
        span = len(sundays)
    columns += ["ተገኝቷል", "ቀርቷል", "ፈቃድ"]
    ws.append(columns)
    _style_header(ws, 3, len(columns))
    for r in rows:
        line = [r["n"], r["student"].full_name]
        line += [""] * span if sheet.is_template else [STATUS_SHORT[c["status"]] for c in r["cells"]]
        line += ["", "", ""] if sheet.is_template else [r["present"], r["absent"], r["permission"]]
        ws.append(line)
    ws.column_dimensions["B"].width = 38
    return wb


def roster_workbook(roster, columns, rows):
    """`columns`: RosterColumn queryset/list (uses .subject_name, .pk).
    `rows`: list of dicts — {"n", "student" (has .full_name), "scores" (dict keyed by
    column.pk), "average", "rank"} — built by the roster view from RosterRow/RosterScore."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "ሮስተር"
    label = " / ".join(p for p in [roster.year, roster.semester] if p) or "ሮስተር"
    ws.append([f"ሮስተር (Roster) — {roster.classroom.name} — {label}"])
    ws.append([])
    header = ["#", "ስም (Name)"] + [c.subject_name for c in columns] + ["አማካኝ", "ደረጃ"]
    ws.append(header)
    _style_header(ws, 3, len(header))
    gold = PatternFill("solid", fgColor="FFF2B3")
    for row in rows:
        line = [row["n"], row["student"].full_name]
        line += [None if row["scores"].get(c.pk) is None else float(row["scores"][c.pk]) for c in columns]
        line += [
            float(row["average"]) if row["average"] is not None else None,
            row["rank"],
        ]
        ws.append(line)
        if row["rank"] and row["rank"] <= 3:
            for cell in ws[ws.max_row]:
                cell.fill = gold
                cell.font = Font(bold=True)
    ws.column_dimensions["B"].width = 38
    return wb